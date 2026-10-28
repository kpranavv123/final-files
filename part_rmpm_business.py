import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ─────────────────────────────────────────────
#  FILE PATHS
# ─────────────────────────────────────────────
PART_INPUT_FILE              = r"C:\Users\SW526XH\Downloads\Go Live-2\Part\Part.tab"
PARTSOURCE_MAKE_INPUT_FILE   = r"C:\Users\SW526XH\Downloads\Go Live-2\PartSource\PartSource_Make.tab"
SOURCE_CONSTRAINT_INPUT_FILE = r"C:\Users\SW526XH\Downloads\Go Live-2\SourceConstraint\SourceConstraint.tab"
OUTPUT_FILE                  = r"C:\Users\SW526XH\Downloads\Go Live-2\Part\Validated_Part_Business.xlsx"


# ─────────────────────────────────────────────
#  Colours / Styles  — matched to existing Business Validator template
# ─────────────────────────────────────────────
RED_FILL        = PatternFill("solid", start_color="FF0000", end_color="FF0000")
HDR_FILL        = PatternFill("solid", start_color="D9E1F2", end_color="D9E1F2")
RULE_FILL       = PatternFill("solid", start_color="E2EFDA", end_color="E2EFDA")
TITLE_FILL      = PatternFill("solid", start_color="BDD7EE", end_color="BDD7EE")
TOTAL_FILL      = PatternFill("solid", start_color="F2F2F2", end_color="F2F2F2")
WHITE_FILL      = PatternFill("solid", start_color="FFFFFF", end_color="FFFFFF")
STATS_FILL      = PatternFill("solid", start_color="EDEDED", end_color="EDEDED")
SUMM_HDR_FILL   = PatternFill("solid", start_color="BDD7EE", end_color="BDD7EE")
SUMM_TITLE_FILL = PatternFill("solid", start_color="BDD7EE", end_color="BDD7EE")

HDR_FONT    = Font(bold=True, name="Arial")
BODY_FONT   = Font(name="Arial", size=10)
ERR_FONT    = Font(name="Arial", size=10, bold=True, color="FFFFFF")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)


# ─────────────────────────────────────────────
#  COLUMN CONFIG
# ─────────────────────────────────────────────
RULE_FIELD = "PROCUREMENTYPE"

# Part extract
PART_MATERIAL_COL = "MATERIALNUMBER"
PART_PLANT_COL    = "PLANT"
PROCUREMENT_COL   = "PROCUREMENTYPE"
PROCUREMENT_VALUE = "E"

# PartSource (make) extract
PS_MATERIAL_COL = "MATERIAL"
PS_PLANT_COL    = "PLANT"

# Source Constraint extract
SC_MATERIAL_COL = "MATERIAL"
SC_PLANT_COL    = "PLANT"
SC_FACTOR_COL   = "CONSTRAINTFACTOR"

# Sub-rule (b) behaviour when an "E" MATERIALNUMBER-PLANT has NO row at all in
# the Source Constraint extract:
#   False -> not flagged (rule only fails when a matching row exists with factor 0)
#   True  -> flagged as an error as well
MISSING_IN_CONSTRAINT_IS_ERROR = False

# Sub-rule ids
SUB_A = "a"
SUB_B = "b"

SUB_RULE_REASON = {
    SUB_A: "a) PROCUREMENTYPE is 'E' but MATERIALNUMBER-PLANT combination is not present "
           "in PartSource (Make) extract (MATERIAL-PLANT)",
    SUB_B: "b) PROCUREMENTYPE is 'E' but CONSTRAINTFACTOR is '0' for the MATERIAL-PLANT "
           "combination in Source Constraint extract",
}

FIELD_ORDER = [RULE_FIELD]


# ══════════════════════════════════════════════
#  Rule Engine
# ══════════════════════════════════════════════
class PartBusinessRuleEngine:
    """
    Rule: PROCUREMENTYPE
      a) If PROCUREMENTYPE = 'E' for a MATERIALNUMBER-PLANT combination in the Part
         extract, that combination must be present in the PartSource (Make) extract
         (MATERIAL, PLANT).
      b) If PROCUREMENTYPE = 'E' for a MATERIALNUMBER-PLANT combination in the Part
         extract, CONSTRAINTFACTOR must not be '0' for that MATERIAL-PLANT
         combination in the Source Constraint extract.
    Both are cross-extract checks, pre-computed over the whole dataframe.
    """

    @staticmethod
    def _clean(series: pd.Series) -> pd.Series:
        return series.fillna("").astype(str).str.strip()

    @staticmethod
    def _parse_number(value):
        s = str(value).strip()
        if s == "" or s.lower() == "nan":
            return None
        try:
            return float(s)
        except ValueError:
            return None

    @staticmethod
    def _missing_cols(df: pd.DataFrame, cols: list) -> list:
        return [c for c in cols if c not in df.columns]

    def _e_rows(self, part_df: pd.DataFrame) -> pd.DataFrame:
        """Part rows with PROCUREMENTYPE = 'E' and non-blank MATERIALNUMBER + PLANT.
        Returns a frame indexed like part_df with cleaned mat / plant columns."""
        proc  = self._clean(part_df[PROCUREMENT_COL]).str.upper()
        mat   = self._clean(part_df[PART_MATERIAL_COL])
        plant = self._clean(part_df[PART_PLANT_COL])
        mask  = (proc == PROCUREMENT_VALUE) & (mat != "") & (plant != "")
        return pd.DataFrame({"mat": mat[mask], "plant": plant[mask]})

    # ── Sub-rule (a) ──
    def check_partsource(self, part_df: pd.DataFrame, ps_df: pd.DataFrame) -> dict:
        """Returns {row_idx: reason} for E rows whose MATERIALNUMBER-PLANT is missing
        from the PartSource (Make) extract."""
        miss_part = self._missing_cols(part_df, [PROCUREMENT_COL, PART_MATERIAL_COL, PART_PLANT_COL])
        miss_ps   = self._missing_cols(ps_df, [PS_MATERIAL_COL, PS_PLANT_COL])
        if miss_part or miss_ps:
            print(f"⚠️   Sub-rule (a) skipped — missing column(s): Part={miss_part}, PartSource={miss_ps}")
            return {}

        ps_keys = set(zip(self._clean(ps_df[PS_MATERIAL_COL]),
                          self._clean(ps_df[PS_PLANT_COL])))

        e_rows = self._e_rows(part_df)
        result = {}
        for idx, r in e_rows.iterrows():
            if (r["mat"], r["plant"]) not in ps_keys:
                result[idx] = (
                    f"PROCUREMENTYPE: 'E' but {PART_MATERIAL_COL} '{r['mat']}' + "
                    f"{PART_PLANT_COL} '{r['plant']}' is not present in PartSource (Make) extract"
                )
        return result

    # ── Sub-rule (b) ──
    def check_constraint_factor(self, part_df: pd.DataFrame, sc_df: pd.DataFrame) -> dict:
        """Returns {row_idx: reason} for E rows whose MATERIAL-PLANT has CONSTRAINTFACTOR = 0
        in the Source Constraint extract."""
        miss_part = self._missing_cols(part_df, [PROCUREMENT_COL, PART_MATERIAL_COL, PART_PLANT_COL])
        miss_sc   = self._missing_cols(sc_df, [SC_MATERIAL_COL, SC_PLANT_COL, SC_FACTOR_COL])
        if miss_part or miss_sc:
            print(f"⚠️   Sub-rule (b) skipped — missing column(s): Part={miss_part}, SourceConstraint={miss_sc}")
            return {}

        sc = pd.DataFrame({
            "mat":    self._clean(sc_df[SC_MATERIAL_COL]),
            "plant":  self._clean(sc_df[SC_PLANT_COL]),
            "factor": self._clean(sc_df[SC_FACTOR_COL]),
        })

        # Per combination: list of raw factors, and whether any of them is zero
        factors_by_key = {}
        zero_keys      = set()
        for mat, plant, factor in zip(sc["mat"], sc["plant"], sc["factor"]):
            key = (mat, plant)
            factors_by_key.setdefault(key, []).append(factor)
            num = self._parse_number(factor)
            if num is not None and num == 0:
                zero_keys.add(key)

        e_rows = self._e_rows(part_df)
        result = {}
        for idx, r in e_rows.iterrows():
            key = (r["mat"], r["plant"])
            if key in zero_keys:
                result[idx] = (
                    f"PROCUREMENTYPE: 'E' but {SC_FACTOR_COL} is '0' for {SC_MATERIAL_COL} "
                    f"'{r['mat']}' + {SC_PLANT_COL} '{r['plant']}' in Source Constraint extract"
                )
            elif MISSING_IN_CONSTRAINT_IS_ERROR and key not in factors_by_key:
                result[idx] = (
                    f"PROCUREMENTYPE: 'E' but {SC_MATERIAL_COL} '{r['mat']}' + "
                    f"{SC_PLANT_COL} '{r['plant']}' is not present in Source Constraint extract"
                )
        return result


# ══════════════════════════════════════════════
#  Validator
# ══════════════════════════════════════════════
class PartBusinessTableValidator:

    def __init__(self, part_path: str, ps_path: str, sc_path: str):
        self.part_path = part_path
        self.ps_path   = ps_path
        self.sc_path   = sc_path
        self.df        = pd.DataFrame()
        self.ps_df     = pd.DataFrame()
        self.sc_df     = pd.DataFrame()
        self.error_map  = {}   # row_idx -> [failed field names]
        self.reason_map = {}   # row_idx -> {field: reason}
        self.sub_map    = {}   # row_idx -> {sub_rule_id: reason}

    @staticmethod
    def _read(path: str) -> pd.DataFrame:
        df = pd.read_csv(path, sep="\t", dtype=str)
        df.columns = [c.strip().upper() for c in df.columns]
        return df

    def load(self):
        self.df    = self._read(self.part_path)
        self.ps_df = self._read(self.ps_path)
        self.sc_df = self._read(self.sc_path)

    def validate(self):
        engine = PartBusinessRuleEngine()

        a_reasons = engine.check_partsource(self.df, self.ps_df)
        b_reasons = engine.check_constraint_factor(self.df, self.sc_df)

        for sub_id, reasons in ((SUB_A, a_reasons), (SUB_B, b_reasons)):
            for idx, reason in reasons.items():
                self.sub_map.setdefault(idx, {})[sub_id] = reason

        for idx, subs in self.sub_map.items():
            self.error_map[idx]  = [RULE_FIELD]
            self.reason_map[idx] = {
                RULE_FIELD: " | ".join(f"[{sid}] {txt}" for sid, txt in subs.items())
            }

    def get_error_series(self) -> pd.Series:
        result = {}
        for idx, col_reason in self.reason_map.items():
            result[idx] = " | ".join(col_reason.values())
        return pd.Series(result, dtype=str)

    def get_field_error_series(self, field_name: str) -> pd.Series:
        result = {}
        for idx, col_reason in self.reason_map.items():
            if field_name in col_reason:
                result[idx] = col_reason[field_name]
        return pd.Series(result, dtype=str)

    def get_errors_by_field(self) -> dict:
        field_errors: dict = {}
        for row_idx, bad_cols in self.error_map.items():
            for col in bad_cols:
                field_errors.setdefault(col, []).append(row_idx)
        return field_errors

    def get_sub_rule_counts(self) -> dict:
        counts = {SUB_A: 0, SUB_B: 0}
        for subs in self.sub_map.values():
            for sid in subs:
                counts[sid] += 1
        return counts


# ══════════════════════════════════════════════
#  Report Writer
# ══════════════════════════════════════════════
class PartBusinessReportWriter:

    SHEET_SUMMARY = "Summary"
    SHEET_RULES   = "Rules"

    RULES_CONTENT = {
        RULE_FIELD: [
            "a) If PROCUREMENTYPE is 'E' for a MATERIALNUMBER-PLANT combination in the "
            "Part extract, that combination must be present in the PartSource (Make) "
            "extract (MATERIAL and PLANT columns).",
            "b) If PROCUREMENTYPE is 'E' for a MATERIALNUMBER-PLANT combination in the "
            "Part extract, CONSTRAINTFACTOR must not be '0' for that MATERIAL-PLANT "
            "combination in the Source Constraint extract.",
        ],
    }

    def __init__(self, validator: PartBusinessTableValidator, output_path: str):
        self.validator   = validator
        self.output_path = output_path

    # ── helpers ──────────────────────────────
    def _write_header(self, ws, columns):
        for c_idx, col_name in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=c_idx, value=col_name)
            if col_name == "ERROR_COLUMNS":
                cell.fill = WHITE_FILL
                cell.font = Font(bold=True, name="Arial", color="000000")
            else:
                cell.fill = HDR_FILL
                cell.font = HDR_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border    = THIN_BORDER

    def _set_widths(self, ws):
        for col in ws.columns:
            max_len = max((len(str(c.value)) if c.value else 0) for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 60)

    def _style_summary_data_row(self, ws, row_num: int, num_cols: int = 7,
                                bold: bool = False, fill: PatternFill = None,
                                italic: bool = False):
        for c in range(1, num_cols + 1):
            cell           = ws.cell(row=row_num, column=c)
            cell.font      = Font(name="Arial", bold=bold, italic=italic, size=10)
            cell.border    = THIN_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if fill:
                cell.fill = fill

    def _highlight_cols_for(self, field_name: str) -> list:
        """Part columns highlighted red when the PROCUREMENTYPE rule fails."""
        return [PROCUREMENT_COL, PART_MATERIAL_COL, PART_PLANT_COL]

    # ══════════════════════════════════════════
    #  Summary sheet  (field row + one sub-row per sub-rule)
    # ══════════════════════════════════════════
    def _write_summary_sheet_into(self, ws, error_map: dict, total_rows: int):

        # ── Row 1 : Title ──
        ws.merge_cells("A1:G1")
        title_cell           = ws.cell(row=1, column=1, value="Part Business Validation Summary")
        title_cell.font      = Font(name="Arial", bold=True, size=14)
        title_cell.fill      = SUMM_TITLE_FILL
        title_cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 26

        # ── Row 2 : Column headers ──
        headers = ["#", "Field Name", "Error Count", "Record Count",
                   "% Health", "% of Error", "Reason / Sub-Category"]
        for c_idx, h in enumerate(headers, start=1):
            cell           = ws.cell(row=2, column=c_idx, value=h)
            cell.fill      = SUMM_HDR_FILL
            cell.font      = Font(name="Arial", bold=True, size=10)
            cell.border    = THIN_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[2].height = 30

        col_error_counts: dict = {}
        for bad_cols in error_map.values():
            for col in bad_cols:
                col_error_counts[col] = col_error_counts.get(col, 0) + 1

        sub_counts = self.validator.get_sub_rule_counts()

        row_num   = 3
        field_num = 1

        for col_name in FIELD_ORDER:
            count      = col_error_counts.get(col_name, 0)
            pct_error  = round((count / total_rows) * 100, 2) if total_rows else 0
            pct_health = round(100 - pct_error, 2)

            ws.cell(row=row_num, column=1, value=field_num)
            ws.cell(row=row_num, column=2, value=col_name)
            ws.cell(row=row_num, column=3, value=count)
            ws.cell(row=row_num, column=4, value=total_rows)
            ws.cell(row=row_num, column=5, value=f"{pct_health}%")
            ws.cell(row=row_num, column=6, value=f"{pct_error}%")
            ws.cell(row=row_num, column=7, value="")
            self._style_summary_data_row(ws, row_num, bold=True, fill=WHITE_FILL)
            row_num += 1

            # Sub-rule rows
            for sid in (SUB_A, SUB_B):
                s_count = sub_counts.get(sid, 0)
                s_err   = round((s_count / total_rows) * 100, 2) if total_rows else 0
                s_hlth  = round(100 - s_err, 2)

                ws.cell(row=row_num, column=1, value="")
                ws.cell(row=row_num, column=2, value=f"   {col_name} - sub-rule {sid})")
                ws.cell(row=row_num, column=3, value=s_count)
                ws.cell(row=row_num, column=4, value=total_rows)
                ws.cell(row=row_num, column=5, value=f"{s_hlth}%")
                ws.cell(row=row_num, column=6, value=f"{s_err}%")
                ws.cell(row=row_num, column=7,
                        value=SUB_RULE_REASON[sid] if s_count > 0 else "")
                self._style_summary_data_row(ws, row_num, fill=WHITE_FILL, italic=True)
                ws.cell(row=row_num, column=2).alignment = Alignment(horizontal="left", vertical="center")
                ws.cell(row=row_num, column=7).alignment = Alignment(
                    horizontal="left", vertical="center", wrap_text=True
                )
                row_num += 1

            field_num += 1

        # ── TOTAL row (unique failing rows for the rule field) ──
        total_errors       = sum(col_error_counts.values())
        total_record_count = total_rows * len(FIELD_ORDER)
        total_pct_error    = round((total_errors / total_record_count) * 100, 2) if total_record_count else 0
        total_pct_health   = round(100 - total_pct_error, 2)

        ws.cell(row=row_num, column=1, value="")
        ws.cell(row=row_num, column=2, value="TOTAL")
        ws.cell(row=row_num, column=3, value=total_errors)
        ws.cell(row=row_num, column=4, value=total_record_count)
        ws.cell(row=row_num, column=5, value=f"{total_pct_health}%")
        ws.cell(row=row_num, column=6, value=f"{total_pct_error}%")
        ws.cell(row=row_num, column=7, value="")

        for c in range(1, 8):
            ws.cell(row=row_num, column=c).font      = Font(name="Arial", bold=True, size=10)
            ws.cell(row=row_num, column=c).fill      = TOTAL_FILL
            ws.cell(row=row_num, column=c).border    = THIN_BORDER
            ws.cell(row=row_num, column=c).alignment = Alignment(horizontal="center", vertical="center")

        row_num += 2   # blank spacer

        # ── Quick-glance stats block ──
        records_with_errors = len(error_map)
        records_passing     = total_rows - records_with_errors

        for label, value in [
            ("Total Records:",       total_rows),
            ("Records with Errors:", records_with_errors),
            ("Records Passing:",     records_passing),
        ]:
            ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=2)
            label_cell           = ws.cell(row=row_num, column=1, value=label)
            label_cell.font      = Font(name="Arial", bold=True, size=10)
            label_cell.fill      = STATS_FILL
            label_cell.border    = THIN_BORDER
            label_cell.alignment = Alignment(horizontal="left", vertical="center")

            value_cell           = ws.cell(row=row_num, column=3, value=value)
            value_cell.font      = Font(name="Arial", size=10)
            value_cell.border    = THIN_BORDER
            value_cell.alignment = Alignment(horizontal="center", vertical="center")

            row_num += 1

        col_widths = [6, 42, 14, 16, 12, 12, 90]
        for c_idx, width in enumerate(col_widths, start=1):
            ws.column_dimensions[get_column_letter(c_idx)].width = width

    # ── Per-field error sheets  (ALL source columns) ──
    def _write_field_error_sheets(self, wb, df: pd.DataFrame, all_cols: list):
        field_errors = self.validator.get_errors_by_field()

        for field_name in FIELD_ORDER:
            if field_name not in field_errors:
                continue

            row_indices = field_errors[field_name]
            sheet_name  = field_name[:31].replace("/", "-").replace("\\", "-").replace("*", "")
            ws          = wb.create_sheet(sheet_name)

            keep_here = all_cols + ["ERROR_COLUMNS"]
            subset    = df.loc[row_indices, keep_here].copy()

            field_err_series        = self.validator.get_field_error_series(field_name)
            subset["ERROR_COLUMNS"] = subset.index.map(
                lambda i: field_err_series.get(i, "")
            )

            # Sort so rows of the same MATERIALNUMBER-PLANT sit together
            sort_cols = [c for c in (PART_MATERIAL_COL, PART_PLANT_COL) if c in subset.columns]
            if sort_cols:
                subset = subset.sort_values(sort_cols, kind="stable", na_position="last")

            self._write_header(ws, subset.columns)
            col_idx_map = {col: i for i, col in enumerate(subset.columns, start=1)}

            for excel_row, (orig_idx, row_data) in enumerate(subset.iterrows(), start=2):
                for c_idx, (col, value) in enumerate(zip(subset.columns, row_data), start=1):
                    cell           = ws.cell(row=excel_row, column=c_idx, value=value)
                    cell.font      = BODY_FONT
                    cell.alignment = Alignment(vertical="center")
                    cell.fill      = WHITE_FILL
                    cell.border    = THIN_BORDER

                # Highlight PROCUREMENTYPE + MATERIALNUMBER + PLANT on failing rows
                row_reasons = self.validator.reason_map.get(orig_idx, {})
                for involved_field in row_reasons:
                    for target_col in self._highlight_cols_for(involved_field):
                        if target_col in col_idx_map:
                            target_cell      = ws.cell(row=excel_row, column=col_idx_map[target_col])
                            target_cell.fill = RED_FILL
                            target_cell.font = ERR_FONT

            self._set_widths(ws)
            ws.freeze_panes = "A2"

            note_row = len(subset) + 3
            ws.cell(
                row=note_row, column=1,
                value=f"Total error rows for '{field_name}': {len(subset)}",
            ).font = Font(name="Arial", italic=True, size=9, bold=True)

    # ── Rules sheet ───────────────────────────
    def _write_rules_sheet(self, wb):
        ws = wb.create_sheet(self.SHEET_RULES)

        ws.merge_cells("A1:C1")
        title_cell           = ws.cell(row=1, column=1, value="Part Table – Business Validation Rules")
        title_cell.font      = Font(name="Arial", bold=True, size=13)
        title_cell.fill      = TITLE_FILL
        title_cell.alignment = Alignment(horizontal="center")
        ws.row_dimensions[1].height = 22

        for c_idx, h in enumerate(["#", "Field", "Rule Description"], start=1):
            cell           = ws.cell(row=3, column=c_idx, value=h)
            cell.fill      = HDR_FILL
            cell.font      = HDR_FONT
            cell.border    = THIN_BORDER
            cell.alignment = Alignment(horizontal="center")

        current_row = 4
        rule_num    = 1

        for field in FIELD_ORDER:
            rules_list = self.RULES_CONTENT.get(field, [])
            num_rules  = len(rules_list)

            for r_idx, rule_text in enumerate(rules_list):
                num_cell           = ws.cell(row=current_row, column=1,
                                             value=rule_num if r_idx == 0 else "")
                num_cell.font      = Font(name="Arial", size=10, bold=(r_idx == 0))
                num_cell.fill      = RULE_FILL
                num_cell.border    = THIN_BORDER
                num_cell.alignment = Alignment(horizontal="center", vertical="center")

                field_cell           = ws.cell(row=current_row, column=2,
                                               value=field if r_idx == 0 else "")
                field_cell.font      = Font(name="Arial", size=10, bold=(r_idx == 0))
                field_cell.fill      = RULE_FILL
                field_cell.border    = THIN_BORDER
                field_cell.alignment = Alignment(vertical="center")

                desc_cell           = ws.cell(row=current_row, column=3, value=rule_text)
                desc_cell.font      = BODY_FONT
                desc_cell.border    = THIN_BORDER
                desc_cell.alignment = Alignment(wrap_text=True, vertical="center")

                current_row += 1

            if num_rules > 1:
                s = current_row - num_rules
                e = current_row - 1
                ws.merge_cells(start_row=s, start_column=1, end_row=e, end_column=1)
                ws.merge_cells(start_row=s, start_column=2, end_row=e, end_column=2)

            rule_num += 1

        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 100

    # ── Main write ────────────────────────────
    def write(self):
        v  = self.validator
        df = v.df.copy()

        all_cols = list(df.columns)

        error_series        = v.get_error_series()
        df["ERROR_COLUMNS"] = df.index.map(
            lambda i: error_series.get(i, "") if i in error_series.index else ""
        )

        wb               = Workbook()
        ws_summary       = wb.active
        ws_summary.title = self.SHEET_SUMMARY
        self._write_summary_sheet_into(ws_summary, v.error_map, total_rows=len(df))

        self._write_rules_sheet(wb)
        self._write_field_error_sheets(wb, df, all_cols)

        wb.save(self.output_path)

        sub_counts = v.get_sub_rule_counts()
        fields_with_errors = [f for f in FIELD_ORDER if f in v.get_errors_by_field()]
        print(f"\n✅  Output saved  → {self.output_path}")
        print(f"   Total rows    : {len(df)}")
        print(f"   Error rows    : {len(v.error_map)}")
        print(f"   Sub-rule (a)  : {sub_counts[SUB_A]}")
        print(f"   Sub-rule (b)  : {sub_counts[SUB_B]}")
        print(f"   Field sheets  : {fields_with_errors}")


# ══════════════════════════════════════════════
#  Orchestrator
# ══════════════════════════════════════════════
class PartBusinessTableProcessor:

    def __init__(self, part_path: str, ps_path: str, sc_path: str, output_path: str):
        self.validator = PartBusinessTableValidator(part_path, ps_path, sc_path)
        self.writer    = PartBusinessReportWriter(self.validator, output_path)

    def run(self):
        print("📂  Loading files …")
        self.validator.load()
        print(f"    Part columns detected             : {list(self.validator.df.columns)}")
        print(f"    PartSource (Make) columns detected: {list(self.validator.ps_df.columns)}")
        print(f"    Source Constraint columns detected: {list(self.validator.sc_df.columns)}")
        print("🔍  Validating business rules …")
        self.validator.validate()
        print("📝  Writing report …")
        self.writer.write()


# ══════════════════════════════════════════════
#  Entry Point
# ══════════════════════════════════════════════
if __name__ == "__main__":
    processor = PartBusinessTableProcessor(
        part_path   = PART_INPUT_FILE,
        ps_path     = PARTSOURCE_MAKE_INPUT_FILE,
        sc_path     = SOURCE_CONSTRAINT_INPUT_FILE,
        output_path = OUTPUT_FILE,
    )
    processor.run()
