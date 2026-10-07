import re
from collections import defaultdict
from datetime import datetime
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ─────────────────────────────────────────────
#  FILE PATHS
# ─────────────────────────────────────────────
PARTSOURCE_BUY_INPUT_FILE = r"C:\Users\SW526XH\Downloads\Go Live-2\PartSource_Buy\PartSource_Buy_2026-07-17-1905.tab"
PARTMASTER_INPUT_FILE     = r"C:\Users\SW526XH\Downloads\Go Live-2\Part\Part.tab"   # <-- update to your Part master file
OUTPUT_FILE                = r"C:\Users\SW526XH\Downloads\Go Live-2\PartSource_Buy\Validated_PartSourceBuy_Business.xlsx"


# ─────────────────────────────────────────────
#  Colours / Styles  — matched to existing Business Validator template
# ─────────────────────────────────────────────
RED_FILL        = PatternFill("solid", start_color="FF0000", end_color="FF0000")
HDR_FILL        = PatternFill("solid", start_color="D9E1F2", end_color="D9E1F2")
RULE_FILL       = PatternFill("solid", start_color="E2EFDA", end_color="E2EFDA")
TITLE_FILL      = PatternFill("solid", start_color="BDD7EE", end_color="BDD7EE")
TOTAL_FILL      = PatternFill("solid", start_color="F2F2F2", end_color="F2F2F2")
WHITE_FILL      = PatternFill("solid", start_color="FFFFFF", end_color="FFFFFF")
STATS_FILL      = PatternFill("solid", start_color="EDEDED", end_color="EDEDED")
SUMM_HDR_FILL   = PatternFill("solid", start_color="BDD7EE", end_color="BDD7EE")
SUMM_TITLE_FILL = PatternFill("solid", start_color="BDD7EE", end_color="BDD7EE")
SUB_FILL        = PatternFill("solid", start_color="FFFFFF", end_color="FFFFFF")

HDR_FONT    = Font(bold=True, name="Arial")
BODY_FONT   = Font(name="Arial", size=10)
ERR_FONT    = Font(name="Arial", size=10, bold=True, color="FFFFFF")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)

# ─────────────────────────────────────────────
#  PART MASTER LOOKUP CONFIG
# ─────────────────────────────────────────────
# PartSource(Buy) columns used by the MATERIALNUMBER rule
PS_MATERIAL_COL    = "MATERIALNUMBER"
PS_PLANT_COL       = "PLANT"
PS_ORDERPOLICY_COL = "ORDERPOLICY"

# Part master columns (adjust here if the Part master uses different names)
PM_MATERIAL_COL = "MATERIALNUMBER"
PM_PLANT_COL    = "PLANT"
PM_PROCTYPE_COL = "PROCUREMENTTYPE"

# ORDERPOLICY (upper-cased)  ->  (display label, required PROCUREMENTTYPE)
ORDERPOLICY_RULES = {
    "BUY":    ("BUY",    "F"),
    "SUBCON": ("SubCon", "30"),
}

# ─────────────────────────────────────────────
#  FIELD ORDER  (only fields with business rules)
# ─────────────────────────────────────────────
FIELD_ORDER = ["MATERIALNUMBER", "MAXIMUMPOQUANTITY", "MINIMUMPOQUANTITY",
               "ROUNDINGVALUE", "PLANNEDELIVERYTIME"]

# Fields that get the simple "must be > 0" check
POSITIVE_FIELDS = ["MAXIMUMPOQUANTITY", "MINIMUMPOQUANTITY",
                   "ROUNDINGVALUE", "PLANNEDELIVERYTIME"]

FIELDS_WITH_SUB_ROWS = {
    "MATERIALNUMBER",
    "MAXIMUMPOQUANTITY",
    "MINIMUMPOQUANTITY",
}

# Single-line reason shown in Summary sheet for each field
FIELD_REASON = {
    "MATERIALNUMBER":     "",   # sub-rows carry reasons
    "MAXIMUMPOQUANTITY":  "",   # sub-rows carry reasons
    "MINIMUMPOQUANTITY":  "",   # sub-rows carry reasons
    "ROUNDINGVALUE":      "ROUNDINGVALUE: Value not greater than 0",
    "PLANNEDELIVERYTIME": "PLANNEDELIVERYTIME: Value not greater than 0",
}

# Columns highlighted red on a field's error sheet (default: the field itself).
# Cross-field rules highlight every column involved.
FIELD_HIGHLIGHT_COLS = {
    "MATERIALNUMBER": [PS_MATERIAL_COL, PS_PLANT_COL, PS_ORDERPOLICY_COL],
}


# ══════════════════════════════════════════════
#  Rule Engine
# ══════════════════════════════════════════════

class PartSourceBuyBusinessRuleEngine:
    """
    Rules:
      MATERIALNUMBER      - For a Part-Site (MATERIALNUMBER + PLANT) with ORDERPOLICY 'BUY',
                            PROCUREMENTTYPE in Part master must be 'F'.
                          - For a Part-Site with ORDERPOLICY 'SubCon',
                            PROCUREMENTTYPE in Part master must be '30'.
      MAXIMUMPOQUANTITY   - must be > 0.
                          - maximum should be greater than or equal to minimum.
      MINIMUMPOQUANTITY   - must be > 0.
                          - minimum should be lesser than or equal to maximum.
      ROUNDINGVALUE       - must be > 0.
      PLANNEDELIVERYTIME  - must be > 0.
    """

    def __init__(self, partmaster_lookup=None):
        # {(material, plant): {procurement types found in Part master}}
        # None -> MATERIALNUMBER rule is skipped
        self.partmaster_lookup = partmaster_lookup

    @staticmethod
    def _parse_number(value):
        s = str(value).strip()
        if s == "" or s.lower() == "nan":
            return None
        try:
            return float(s)
        except ValueError:
            return None

    @staticmethod
    def _clean(value) -> str:
        return "" if pd.isna(value) else str(value).strip()

    def _validate_positive_field(self, row, field_name: str) -> list:
        """Returns list of reason strings if field_name is present but not > 0."""
        reasons = []

        raw = row.get(field_name, "")
        num = self._parse_number(raw)

        if num is not None and num <= 0:
            reasons.append(f"{field_name}: '{str(raw).strip()}' is not greater than 0")

        return reasons

    def _validate_min_max_relation(self, row) -> dict:
        """
        Validates relation between MINIMUMPOQUANTITY and MAXIMUMPOQUANTITY.

        If minimum > maximum:
          - MINIMUMPOQUANTITY fails because minimum should be <= maximum
          - MAXIMUMPOQUANTITY fails because maximum should be >= minimum
        """
        reasons = {
            "MAXIMUMPOQUANTITY": [],
            "MINIMUMPOQUANTITY": [],
        }

        min_raw = row.get("MINIMUMPOQUANTITY", "")
        max_raw = row.get("MAXIMUMPOQUANTITY", "")

        min_num = self._parse_number(min_raw)
        max_num = self._parse_number(max_raw)

        # Apply relation check only when both values are valid numbers
        if min_num is not None and max_num is not None:
            if min_num > max_num:
                reasons["MINIMUMPOQUANTITY"].append(
                    f"MINIMUMPOQUANTITY: Minimum '{str(min_raw).strip()}' is greater than MAXIMUMPOQUANTITY '{str(max_raw).strip()}'"
                )
                reasons["MAXIMUMPOQUANTITY"].append(
                    f"MAXIMUMPOQUANTITY: Maximum '{str(max_raw).strip()}' is less than MINIMUMPOQUANTITY '{str(min_raw).strip()}'"
                )

        return reasons

    def _validate_procurement_type(self, row) -> list:
        """
        Part-Site (MATERIALNUMBER + PLANT) lookup against Part master.
          ORDERPOLICY BUY    -> PROCUREMENTTYPE must be 'F'
          ORDERPOLICY SubCon -> PROCUREMENTTYPE must be '30'
        Other order policies, and rows with a blank material/plant, are skipped.
        """
        if self.partmaster_lookup is None:
            return []

        policy_key = self._clean(row.get(PS_ORDERPOLICY_COL, "")).upper()
        if policy_key not in ORDERPOLICY_RULES:
            return []

        material = self._clean(row.get(PS_MATERIAL_COL, ""))
        plant    = self._clean(row.get(PS_PLANT_COL, ""))
        if not material or not plant:
            return []

        label, expected = ORDERPOLICY_RULES[policy_key]
        found = self.partmaster_lookup.get((material, plant))

        if found is None:
            return [
                f"MATERIALNUMBER: Part-Site '{material}-{plant}' with ORDERPOLICY '{label}' "
                f"not found in Part master (expected PROCUREMENTTYPE '{expected}')"
            ]

        if found != {expected}:
            actual = ", ".join(sorted(v if v else "blank" for v in found))
            return [
                f"MATERIALNUMBER: Part-Site '{material}-{plant}' with ORDERPOLICY '{label}' "
                f"has PROCUREMENTTYPE '{actual}' in Part master (expected '{expected}')"
            ]

        return []

    def validate_row(self, row) -> dict:
        """Returns {field_name: [reasons]} for whichever fields failed on this row."""
        reasons = {}

        # Rule 1: Part-Site ORDERPOLICY vs Part master PROCUREMENTTYPE
        proc_reasons = self._validate_procurement_type(row)
        if proc_reasons:
            reasons["MATERIALNUMBER"] = proc_reasons

        # Rule 2: positive check for the numeric fields
        for field_name in POSITIVE_FIELDS:
            field_reasons = self._validate_positive_field(row, field_name)
            if field_reasons:
                reasons[field_name] = field_reasons

        # Rule 3: min/max relationship check
        relation_reasons = self._validate_min_max_relation(row)

        for field_name, field_reasons in relation_reasons.items():
            if field_reasons:
                reasons.setdefault(field_name, []).extend(field_reasons)

        return reasons


# ══════════════════════════════════════════════
#  Validator
# ══════════════════════════════════════════════
class PartSourceBuyBusinessTableValidator:

    def __init__(self, ps_buy_path: str, partmaster_path: str):
        self.ps_buy_path     = ps_buy_path
        self.partmaster_path = partmaster_path
        self.df              = pd.DataFrame()
        self.pm_df           = pd.DataFrame()
        self.partmaster_lookup = None
        self.error_map       = {}   # row_idx -> [failed field names]
        self.reason_map      = {}   # row_idx -> {field: reason}

    def load(self):
        self.df = pd.read_csv(self.ps_buy_path, sep="\t", dtype=str)
        self.df.columns = [c.strip().upper() for c in self.df.columns]

        self.pm_df = pd.read_csv(self.partmaster_path, sep="\t", dtype=str)
        self.pm_df.columns = [c.strip().upper() for c in self.pm_df.columns]

        self._build_partmaster_lookup()

    def _build_partmaster_lookup(self):
        """Builds {(material, plant): {procurement types}} from the Part master.
        Leaves the lookup as None (rule skipped) if required columns are missing."""
        ps_missing = [c for c in (PS_MATERIAL_COL, PS_PLANT_COL, PS_ORDERPOLICY_COL)
                      if c not in self.df.columns]
        pm_missing = [c for c in (PM_MATERIAL_COL, PM_PLANT_COL, PM_PROCTYPE_COL)
                      if c not in self.pm_df.columns]

        if ps_missing or pm_missing:
            if ps_missing:
                print(f"⚠️   MATERIALNUMBER rule skipped — PartSource(Buy) column(s) not found: {ps_missing}")
            if pm_missing:
                print(f"⚠️   MATERIALNUMBER rule skipped — Part master column(s) not found: {pm_missing}")
            self.partmaster_lookup = None
            return

        pm = self.pm_df[[PM_MATERIAL_COL, PM_PLANT_COL, PM_PROCTYPE_COL]].fillna("")
        mats   = pm[PM_MATERIAL_COL].astype(str).str.strip()
        plants = pm[PM_PLANT_COL].astype(str).str.strip()
        ptypes = pm[PM_PROCTYPE_COL].astype(str).str.strip()

        lookup = defaultdict(set)
        for m, p, t in zip(mats, plants, ptypes):
            if m and p:
                lookup[(m, p)].add(t)

        self.partmaster_lookup = dict(lookup)

    def validate(self):
        engine = PartSourceBuyBusinessRuleEngine(self.partmaster_lookup)

        for idx, row in self.df.iterrows():
            try:
                reasons = engine.validate_row(row)
            except Exception:
                reasons = {}

            if reasons:
                self.error_map[idx]  = list(reasons.keys())
                self.reason_map[idx] = reasons

    def get_error_series(self) -> pd.Series:
        result = {}
        for idx, col_reason in self.reason_map.items():
            all_reasons = []

            for reason_value in col_reason.values():
                if isinstance(reason_value, list):
                    all_reasons.extend(reason_value)
                else:
                    all_reasons.append(reason_value)

            result[idx] = " | ".join(all_reasons)

        return pd.Series(result, dtype=str)

    def get_field_error_series(self, field_name: str) -> pd.Series:
        result = {}
        for idx, col_reason in self.reason_map.items():
            if field_name in col_reason:
                reason_value = col_reason[field_name]

                if isinstance(reason_value, list):
                    result[idx] = " | ".join(reason_value)
                else:
                    result[idx] = reason_value

        return pd.Series(result, dtype=str)

    def get_errors_by_field(self) -> dict:
        field_errors: dict = {}
        for row_idx, bad_cols in self.error_map.items():
            for col in bad_cols:
                field_errors.setdefault(col, []).append(row_idx)
        return field_errors

    def get_quantity_error_subcounts(self, field_name: str) -> dict:
        counts = {
            "not_greater_than_zero": 0,
            "min_max_relation": 0,
        }

        for idx, col_reason in self.reason_map.items():
            reason_value = col_reason.get(field_name, "")

            if not reason_value:
                continue

            if not isinstance(reason_value, list):
                reason_list = [reason_value]
            else:
                reason_list = reason_value

            for reason in reason_list:
                reason_lower = str(reason).lower()

                if "not greater than 0" in reason_lower:
                    counts["not_greater_than_zero"] += 1
                elif "maximumpoquantity" in reason_lower and "minimumpoquantity" in reason_lower:
                    counts["min_max_relation"] += 1

        return counts

    def get_material_error_subcounts(self) -> dict:
        """Counts MATERIALNUMBER failures by ORDERPOLICY (BUY vs SubCon)."""
        counts = {"BUY": 0, "SUBCON": 0}

        for idx, col_reason in self.reason_map.items():
            reason_value = col_reason.get("MATERIALNUMBER", "")

            if not reason_value:
                continue

            reason_list = reason_value if isinstance(reason_value, list) else [reason_value]

            for reason in reason_list:
                reason_lower = str(reason).lower()

                if "orderpolicy 'buy'" in reason_lower:
                    counts["BUY"] += 1
                elif "orderpolicy 'subcon'" in reason_lower:
                    counts["SUBCON"] += 1

        return counts


# ══════════════════════════════════════════════
#  Report Writer
# ══════════════════════════════════════════════
class PartSourceBuyBusinessReportWriter:

    SHEET_SUMMARY = "Summary"
    SHEET_RULES   = "Rules"

    RULES_CONTENT = {
        "MATERIALNUMBER": [
            "For a Part-Site (MATERIALNUMBER and PLANT) combination with \"BUY\" ORDERPOLICY, "
            "the part-site's PROCUREMENTTYPE should be maintained as \"F\" in the Part master file.",
            "For a Part-Site combination with \"SubCon\" ORDERPOLICY, "
            "the part-site's PROCUREMENTTYPE should be maintained as \"30\" in the Part master file.",
        ],
        "MAXIMUMPOQUANTITY": [
            "MAXIMUMPOQUANTITY must be greater than 0.",
            "Maximum should be greater than or equal to minimum.",
        ],
        "MINIMUMPOQUANTITY": [
            "MINIMUMPOQUANTITY must be greater than 0.",
            "Minimum should be lesser than or equal to maximum.",
        ],
        "ROUNDINGVALUE": [
            "ROUNDINGVALUE must be greater than 0.",
        ],
        "PLANNEDELIVERYTIME": [
            "PLANNEDELIVERYTIME must be greater than 0.",
        ],
    }

    def __init__(self, validator: PartSourceBuyBusinessTableValidator, output_path: str):
        self.validator   = validator
        self.output_path = output_path

    # ── helpers ──────────────────────────────
    def _write_header(self, ws, columns):
        for c_idx, col_name in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=c_idx, value=col_name)
            if col_name == "ERROR_COLUMNS":
                cell.fill = WHITE_FILL
                cell.font = Font(bold=True, name="Arial", color="000000")
            else:
                cell.fill = HDR_FILL
                cell.font = HDR_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border    = THIN_BORDER

    def _set_widths(self, ws):
        for col in ws.columns:
            max_len = max((len(str(c.value)) if c.value else 0) for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 60)

    def _style_summary_data_row(self, ws, row_num: int, num_cols: int = 7,
                                bold: bool = False, fill: PatternFill = None,
                                italic: bool = False):
        for c in range(1, num_cols + 1):
            cell           = ws.cell(row=row_num, column=c)
            cell.font      = Font(name="Arial", bold=bold, italic=italic, size=10)
            cell.border    = THIN_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if fill:
                cell.fill = fill

    def _write_sub_rows(self, ws, row_num: int, sub_definitions: list, total_rows: int) -> int:
        """
        Writes the given list of (label, count, reason) sub-rows starting at row_num,
        directly beneath the parent field's row. Returns the next free row_num.
        """
        for sub_label, sub_count, sub_reason in sub_definitions:
            sub_pct_err    = round((sub_count / total_rows) * 100, 2) if total_rows else 0
            sub_pct_health = round(100 - sub_pct_err, 2)

            ws.cell(row=row_num, column=1, value="")
            ws.cell(row=row_num, column=2, value=sub_label)
            ws.cell(row=row_num, column=3, value=sub_count)
            ws.cell(row=row_num, column=4, value=total_rows)
            ws.cell(row=row_num, column=5, value=f"{sub_pct_health}%")
            ws.cell(row=row_num, column=6, value=f"{sub_pct_err}%")
            ws.cell(row=row_num, column=7, value=sub_reason)

            self._style_summary_data_row(ws, row_num, fill=SUB_FILL, italic=True)
            ws.cell(row=row_num, column=2).alignment = Alignment(
                horizontal="left", vertical="center", indent=1
            )
            ws.cell(row=row_num, column=7).alignment = Alignment(
                horizontal="left", vertical="center", wrap_text=True
            )

            row_num += 1

        return row_num

    # ══════════════════════════════════════════
    #  Summary sheet
    # ══════════════════════════════════════════
    def _write_summary_sheet_into(self, ws, error_map: dict, total_rows: int):

        # ── Row 1 : Title ──
        ws.merge_cells("A1:G1")
        title_cell           = ws.cell(row=1, column=1, value="PartSource(Buy) Business Validation Summary")
        title_cell.font      = Font(name="Arial", bold=True, size=14)
        title_cell.fill      = SUMM_TITLE_FILL
        title_cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 26

        # ── Row 2 : Column headers ──
        headers = ["#", "Field Name", "Error Count", "Record Count",
                   "% Health", "% of Error", "Reason"]
        for c_idx, h in enumerate(headers, start=1):
            cell           = ws.cell(row=2, column=c_idx, value=h)
            cell.fill      = SUMM_HDR_FILL
            cell.font      = Font(name="Arial", bold=True, size=10)
            cell.border    = THIN_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[2].height = 30

        # ── Build per-field error counts ──
        col_error_counts: dict = {}
        for bad_cols in error_map.values():
            for col in bad_cols:
                col_error_counts[col] = col_error_counts.get(col, 0) + 1

        max_qty_sub = self.validator.get_quantity_error_subcounts("MAXIMUMPOQUANTITY")
        min_qty_sub = self.validator.get_quantity_error_subcounts("MINIMUMPOQUANTITY")
        material_sub = self.validator.get_material_error_subcounts()

        row_num   = 3
        field_num = 1

        for col_name in FIELD_ORDER:
            count      = col_error_counts.get(col_name, 0)
            has_errors = count > 0

            pct_error  = round((count / total_rows) * 100, 2) if total_rows else 0
            pct_health = round(100 - pct_error, 2)

            if col_name in FIELDS_WITH_SUB_ROWS:
                reason_text = ""
            elif has_errors:
                reason_text = FIELD_REASON.get(col_name, "")
            else:
                reason_text = ""

            # ── Main field row ──
            ws.cell(row=row_num, column=1, value=field_num)
            ws.cell(row=row_num, column=2, value=col_name)
            ws.cell(row=row_num, column=3, value=count)
            ws.cell(row=row_num, column=4, value=total_rows)
            ws.cell(row=row_num, column=5, value=f"{pct_health}%")
            ws.cell(row=row_num, column=6, value=f"{pct_error}%")
            ws.cell(row=row_num, column=7, value=reason_text)

            self._style_summary_data_row(ws, row_num, fill=WHITE_FILL)
            ws.cell(row=row_num, column=7).alignment = Alignment(
                horizontal="left", vertical="center", wrap_text=True
            )
            row_num += 1

            # ── Sub-rows immediately beneath their own parent field ──
            if col_name == "MATERIALNUMBER" and has_errors:
                sub_definitions = [
                    (
                        "  ↳ BUY: PROCUREMENTTYPE not 'F'",
                        material_sub["BUY"],
                        "MATERIALNUMBER: Part-Site with ORDERPOLICY 'BUY' should have PROCUREMENTTYPE 'F' in Part master",
                    ),
                    (
                        "  ↳ SubCon: PROCUREMENTTYPE not '30'",
                        material_sub["SUBCON"],
                        "MATERIALNUMBER: Part-Site with ORDERPOLICY 'SubCon' should have PROCUREMENTTYPE '30' in Part master",
                    ),
                ]
                row_num = self._write_sub_rows(ws, row_num, sub_definitions, total_rows)

            elif col_name == "MAXIMUMPOQUANTITY" and has_errors:
                sub_definitions = [
                    (
                        "  ↳ Not greater than 0",
                        max_qty_sub["not_greater_than_zero"],
                        "MAXIMUMPOQUANTITY: Value not greater than 0",
                    ),
                    (
                        "  ↳ Maximum less than Minimum",
                        max_qty_sub["min_max_relation"],
                        "MAXIMUMPOQUANTITY: Maximum should be greater than or equal to minimum",
                    ),
                ]
                row_num = self._write_sub_rows(ws, row_num, sub_definitions, total_rows)

            elif col_name == "MINIMUMPOQUANTITY" and has_errors:
                sub_definitions = [
                    (
                        "  ↳ Not greater than 0",
                        min_qty_sub["not_greater_than_zero"],
                        "MINIMUMPOQUANTITY: Value not greater than 0",
                    ),
                    (
                        "  ↳ Minimum greater than Maximum",
                        min_qty_sub["min_max_relation"],
                        "MINIMUMPOQUANTITY: Minimum should be lesser than or equal to maximum",
                    ),
                ]
                row_num = self._write_sub_rows(ws, row_num, sub_definitions, total_rows)

            field_num += 1

        # ── TOTAL row ──
        total_errors        = sum(col_error_counts.values())
        total_record_count  = total_rows * len(FIELD_ORDER)
        total_pct_error     = round((total_errors / total_record_count) * 100, 2) if total_record_count else 0
        total_pct_health    = round(100 - total_pct_error, 2)

        ws.cell(row=row_num, column=1, value="")
        ws.cell(row=row_num, column=2, value="TOTAL")
        ws.cell(row=row_num, column=3, value=total_errors)
        ws.cell(row=row_num, column=4, value=total_record_count)
        ws.cell(row=row_num, column=5, value=f"{total_pct_health}%")
        ws.cell(row=row_num, column=6, value=f"{total_pct_error}%")
        ws.cell(row=row_num, column=7, value="")

        for c in range(1, 8):
            ws.cell(row=row_num, column=c).font      = Font(name="Arial", bold=True, size=10)
            ws.cell(row=row_num, column=c).fill      = TOTAL_FILL
            ws.cell(row=row_num, column=c).border    = THIN_BORDER
            ws.cell(row=row_num, column=c).alignment = Alignment(horizontal="center", vertical="center")

        row_num += 2   # blank spacer

        # ── Quick-glance stats block ──
        records_with_errors = len(error_map)
        records_passing     = total_rows - records_with_errors

        for label, value in [
            ("Total Records:",       total_rows),
            ("Records with Errors:", records_with_errors),
            ("Records Passing:",     records_passing),
        ]:
            ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=2)
            label_cell           = ws.cell(row=row_num, column=1, value=label)
            label_cell.font      = Font(name="Arial", bold=True, size=10)
            label_cell.fill      = STATS_FILL
            label_cell.border    = THIN_BORDER
            label_cell.alignment = Alignment(horizontal="left", vertical="center")

            value_cell           = ws.cell(row=row_num, column=3, value=value)
            value_cell.font      = Font(name="Arial", size=10)
            value_cell.border    = THIN_BORDER
            value_cell.alignment = Alignment(horizontal="center", vertical="center")

            row_num += 1

        # ── Column widths ──
        col_widths = [6, 42, 14, 16, 12, 12, 70]
        for c_idx, width in enumerate(col_widths, start=1):
            ws.column_dimensions[get_column_letter(c_idx)].width = width

    # ── Per-field error sheets  (ALL source columns) ──
    def _write_field_error_sheets(self, wb, df: pd.DataFrame, all_cols: list):
        field_errors = self.validator.get_errors_by_field()

        for field_name in FIELD_ORDER:
            if field_name not in field_errors:
                continue

            row_indices = field_errors[field_name]
            sheet_name  = field_name[:31].replace("/", "-").replace("\\", "-").replace("*", "")
            ws          = wb.create_sheet(sheet_name)

            keep_here = all_cols + ["ERROR_COLUMNS"]
            subset    = df.loc[row_indices, keep_here].copy()

            field_err_series        = self.validator.get_field_error_series(field_name)
            subset["ERROR_COLUMNS"] = subset.index.map(
                lambda i: field_err_series.get(i, "")
            )

            self._write_header(ws, subset.columns)
            col_idx_map = {col: i for i, col in enumerate(subset.columns, start=1)}

            # Columns to highlight on this sheet: every column involved in the rule
            # (MATERIALNUMBER + PLANT + ORDERPOLICY for the Part master check),
            # otherwise just the field itself.
            highlight_cols = FIELD_HIGHLIGHT_COLS.get(field_name, [field_name])

            for excel_row, (orig_idx, row_data) in enumerate(subset.iterrows(), start=2):
                for c_idx, (col, value) in enumerate(zip(subset.columns, row_data), start=1):
                    cell           = ws.cell(row=excel_row, column=c_idx, value=value)
                    cell.font      = BODY_FONT
                    cell.alignment = Alignment(vertical="center")
                    cell.fill      = WHITE_FILL
                    cell.border    = THIN_BORDER

                for target_col in highlight_cols:
                    if target_col in col_idx_map:
                        target_cell      = ws.cell(row=excel_row, column=col_idx_map[target_col])
                        target_cell.fill = RED_FILL
                        target_cell.font = ERR_FONT

            self._set_widths(ws)
            ws.freeze_panes = "A2"

            note_row = len(subset) + 3
            ws.cell(
                row=note_row, column=1,
                value=f"Total error rows for '{field_name}': {len(subset)}",
            ).font = Font(name="Arial", italic=True, size=9, bold=True)

    # ── Rules sheet ───────────────────────────
    def _write_rules_sheet(self, wb):
        ws = wb.create_sheet(self.SHEET_RULES)

        ws.merge_cells("A1:C1")
        title_cell           = ws.cell(row=1, column=1, value="PartSource(Buy) Table – Business Validation Rules")
        title_cell.font      = Font(name="Arial", bold=True, size=13)
        title_cell.fill      = TITLE_FILL
        title_cell.alignment = Alignment(horizontal="center")
        ws.row_dimensions[1].height = 22

        for c_idx, h in enumerate(["#", "Field", "Rule Description"], start=1):
            cell           = ws.cell(row=3, column=c_idx, value=h)
            cell.fill      = HDR_FILL
            cell.font      = HDR_FONT
            cell.border    = THIN_BORDER
            cell.alignment = Alignment(horizontal="center")

        current_row = 4
        rule_num    = 1

        for field in FIELD_ORDER:
            rules_list = self.RULES_CONTENT.get(field, [])
            num_rules  = len(rules_list)

            for r_idx, rule_text in enumerate(rules_list):
                num_cell           = ws.cell(row=current_row, column=1,
                                             value=rule_num if r_idx == 0 else "")
                num_cell.font      = Font(name="Arial", size=10, bold=(r_idx == 0))
                num_cell.fill      = RULE_FILL
                num_cell.border    = THIN_BORDER
                num_cell.alignment = Alignment(horizontal="center", vertical="center")

                field_cell           = ws.cell(row=current_row, column=2,
                                               value=field if r_idx == 0 else "")
                field_cell.font      = Font(name="Arial", size=10, bold=(r_idx == 0))
                field_cell.fill      = RULE_FILL
                field_cell.border    = THIN_BORDER
                field_cell.alignment = Alignment(vertical="center")

                desc_cell           = ws.cell(row=current_row, column=3, value=rule_text)
                desc_cell.font      = BODY_FONT
                desc_cell.border    = THIN_BORDER
                desc_cell.alignment = Alignment(wrap_text=True, vertical="center")

                current_row += 1

            if num_rules > 1:
                s = current_row - num_rules
                e = current_row - 1
                ws.merge_cells(start_row=s, start_column=1, end_row=e, end_column=1)
                ws.merge_cells(start_row=s, start_column=2, end_row=e, end_column=2)

            rule_num += 1

        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 45
        ws.column_dimensions["C"].width = 75

    # ── Main write ────────────────────────────
    def write(self):
        v  = self.validator
        df = v.df.copy()

        # All source columns, in original order — used for per-field error sheets
        all_cols = list(df.columns)

        error_series        = v.get_error_series()
        df["ERROR_COLUMNS"] = df.index.map(
            lambda i: error_series.get(i, "") if i in error_series.index else ""
        )

        wb               = Workbook()
        ws_summary       = wb.active
        ws_summary.title = self.SHEET_SUMMARY
        self._write_summary_sheet_into(ws_summary, v.error_map, total_rows=len(df))

        self._write_rules_sheet(wb)
        self._write_field_error_sheets(wb, df, all_cols)

        wb.save(self.output_path)

        fields_with_errors = [f for f in FIELD_ORDER if f in v.get_errors_by_field()]
        print(f"\n✅  Output saved  → {self.output_path}")
        print(f"   Total rows    : {len(df)}")
        print(f"   Error rows    : {len(v.error_map)}")
        print(f"   Field sheets  : {fields_with_errors}")


# ══════════════════════════════════════════════
#  Orchestrator
# ══════════════════════════════════════════════
class PartSourceBuyBusinessTableProcessor:

    def __init__(self, ps_buy_path: str, partmaster_path: str, output_path: str):
        self.validator = PartSourceBuyBusinessTableValidator(ps_buy_path, partmaster_path)
        self.writer    = PartSourceBuyBusinessReportWriter(self.validator, output_path)

    def run(self):
        print("📂  Loading files …")
        self.validator.load()
        print(f"    PartSource(Buy) columns detected : {list(self.validator.df.columns)}")
        print(f"    Part master columns detected     : {list(self.validator.pm_df.columns)}")
        print("🔍  Validating business rules …")
        self.validator.validate()
        print("📝  Writing report …")
        self.writer.write()


# ══════════════════════════════════════════════
#  Entry Point
# ══════════════════════════════════════════════
if __name__ == "__main__":
    processor = PartSourceBuyBusinessTableProcessor(
        ps_buy_path     = PARTSOURCE_BUY_INPUT_FILE,
        partmaster_path = PARTMASTER_INPUT_FILE,
        output_path     = OUTPUT_FILE,
    )
    processor.run()
